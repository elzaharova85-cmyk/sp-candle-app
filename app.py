import streamlit as st
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter
import os, io, zipfile, tempfile, json
from datetime import datetime

# ─────────────────────────────────────────────
#  СЕБЕСТОИМОСТЬ  (по типу товара)
# ─────────────────────────────────────────────
DEFAULT_SEBES = {
    'автодиффузор 7мл':            66,
    'диффузор-спрей 100мл':       233.5,
    'диффузор 100мл':              327,
    'диффузор 50мл':               198,
    'диффузор sign':              1629,
    'диффузор-рефил 5л':           660,
    'диффузор-рефил 500мл':       1288,
    'диффузор-рефил 250мл':        684,
    'диффузор 30мл':               288,
    'срезанный флакон 30мл':       288,
    'тестер шайба':                 31,
    'свеча 320г':                 1334,
    'свеча 220г':                  372,
    'свеча 100г':                  174,
    'свеча pur':                   174,
    'крем для рук дой-пак 1 л':    793,
    'крем для рук 1л':             793,
    'крем для рук 45мл':           191,
    'крем для рук 250мл':          290,
    'мыло для рук 250мл':          201,
    'мыло для рук 500мл':          300,
    'мыло для рук 5 л':           1850,
    'мыло для рук 5л':            1850,
    'набор hand care':             875,
    'набор свечей candle trio':    498,
    'набор: свеча220':            1043,
    'подарочная коробка бордовая': 245,
    'фибровые палочки 3/20':        30,
    'фибровые палочки 4/22':        34,
    'фибровые палочки 5/20':        30,
    'спички каминные':              32.5,
    'стакан agger 350мл':         1260,
    'стакан agger 230мл':         1050,
    'подвес':                      280,
}

SEBES_FILE = os.path.join(os.path.dirname(__file__), 'sebes.json')

def load_sebes() -> dict:
    if os.path.exists(SEBES_FILE):
        try:
            with open(SEBES_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            pass
    return dict(DEFAULT_SEBES)

def save_sebes(d: dict):
    with open(SEBES_FILE, 'w', encoding='utf-8') as f:
        json.dump(d, f, ensure_ascii=False, indent=2)

def get_sebes(name: str, sebes_dict: dict = None) -> float | None:
    if not isinstance(name, str):
        return None
    d = sebes_dict if sebes_dict is not None else load_sebes()
    n = name.lower()
    for key, val in d.items():
        if key in n:
            return val
    return None

def _to_float(v) -> float:
    try:
        f = float(v)
        return 0.0 if (f != f) else f  # nan → 0
    except (TypeError, ValueError):
        return 0.0


# ─────────────────────────────────────────────
#  КАТЕГОРИЯ ТОВАРА (для группировки)
# ─────────────────────────────────────────────

# Подстроки, которые относят товар к доставке
DELIVERY_KEYWORDS = [
    'сдэк', 'доставка', 'самовывоз', 'курьер', 'пункт выдачи', 'зарубежье',
    'ближнее зарубежье',
]
# Подстроки для основных товаров (типы продукции)
MAIN_KEYWORDS = [
    'диффузор', 'свеча', 'крем для рук', 'мыло для рук', 'набор', 'шампунь',
    'гель для душа', 'автодиффузор', 'рефил', 'срезанный флакон', 'тестер',
]

def classify_item(name: str) -> str:
    """Возвращает 'основные', 'доставка' или 'сопутствующие'."""
    n = name.lower().strip()
    for kw in DELIVERY_KEYWORDS:
        if kw in n:
            return 'доставка'
    for kw in MAIN_KEYWORDS:
        if kw in n:
            return 'основные'
    return 'сопутствующие'


# ─────────────────────────────────────────────
#  ПАРСИНГ ФАЙЛОВ
# ─────────────────────────────────────────────

def parse_1c_file(file_obj, channel_name: str, sebes_dict: dict) -> pd.DataFrame:
    import xlrd
    with tempfile.NamedTemporaryFile(suffix='.xls', delete=False) as tmp:
        tmp.write(file_obj.read())
        tmp_path = tmp.name

    df_raw = pd.read_excel(tmp_path, engine='xlrd', header=None)
    os.unlink(tmp_path)

    # Ищем строку с "Наименование"
    header_row = None
    for i, row in df_raw.iterrows():
        vals = [str(v).strip() for v in row.values if pd.notna(v)]
        if any('наименование' in v.lower() for v in vals):
            header_row = i
            break

    data_start = (header_row + 2) if header_row is not None else 13

    rows = []
    for i in range(data_start, len(df_raw)):
        row = df_raw.iloc[i]
        name      = row.iloc[3]  if len(row) > 3  else None
        qty       = row.iloc[5]  if len(row) > 5  else None
        price     = row.iloc[6]  if len(row) > 6  else None
        sebes_1c  = row.iloc[7]  if len(row) > 7  else None
        sum_sales = row.iloc[9]  if len(row) > 9  else None
        sum_sebes = row.iloc[10] if len(row) > 10 else None
        profit    = row.iloc[20] if len(row) > 20 else None

        if not isinstance(name, str) or len(name.strip()) < 2:
            continue
        if name.strip().lower() in ('наименование', 'итого', ''):
            continue
        try:
            qty_f = float(qty)
        except (TypeError, ValueError):
            continue
        if qty_f == 0:
            continue

        rows.append({
            'Канал':        channel_name,
            'Наименование': name.strip(),
            'Кол-во':       qty_f,
            'Цена':         _to_float(price),
            'Себес_1С':     _to_float(sebes_1c),
            'Себес_авто':   get_sebes(name, sebes_dict),
            'Сумма продаж': _to_float(sum_sales),
            'Сумма себест': _to_float(sum_sebes),
            'Прибыль':      _to_float(profit),
        })

    return pd.DataFrame(rows)


def parse_ozon_file(file_obj) -> dict:
    raw = file_obj.read()
    xl = pd.ExcelFile(io.BytesIO(raw))
    result = {'_sheet_names': xl.sheet_names}

    for sh in xl.sheet_names:
        sh_low = sh.strip().lower()
        if 'сводк' in sh_low or sh_low in ('сводка', 'summary'):
            result['сводка'] = pd.read_excel(io.BytesIO(raw), sheet_name=sh, header=None)
        elif 'детализ' in sh_low or sh_low in ('детализация', 'detail'):
            result['детализация'] = pd.read_excel(io.BytesIO(raw), sheet_name=sh)
        elif 'типам' in sh_low or ('тип' in sh_low and 'товар' in sh_low):
            result['по_типам'] = pd.read_excel(io.BytesIO(raw), sheet_name=sh, header=None)

    # Если не нашли по имени — берём по позиции
    if 'сводка' not in result and len(xl.sheet_names) >= 1:
        result['сводка'] = pd.read_excel(io.BytesIO(raw), sheet_name=xl.sheet_names[0], header=None)
    if 'детализация' not in result and len(xl.sheet_names) >= 2:
        result['детализация'] = pd.read_excel(io.BytesIO(raw), sheet_name=xl.sheet_names[1])
    if 'по_типам' not in result and len(xl.sheet_names) >= 3:
        result['по_типам'] = pd.read_excel(io.BytesIO(raw), sheet_name=xl.sheet_names[2], header=None)

    return result


def parse_wb_files(zip_files, date_from=None, date_to=None) -> dict:
    all_dfs = []
    for zf in zip_files:
        zf.seek(0)
        with zipfile.ZipFile(zf) as z:
            for name in z.namelist():
                if name.endswith('.xlsx'):
                    with z.open(name) as f:
                        df = pd.read_excel(f, sheet_name='Sheet1')
                        all_dfs.append(df)
    if not all_dfs:
        return {}

    combined = pd.concat(all_dfs, ignore_index=True)

    if date_from or date_to:
        date_col = None
        for cname in combined.columns:
            if 'дата продажи' in str(cname).lower() or 'дата заказа' in str(cname).lower():
                date_col = cname
                break
        if date_col:
            combined[date_col] = pd.to_datetime(combined[date_col], errors='coerce')
            if date_from:
                combined = combined[combined[date_col].dt.date >= date_from]
            if date_to:
                combined = combined[combined[date_col].dt.date <= date_to]

    return {'детализация': combined}


# ─────────────────────────────────────────────
#  СТИЛИ EXCEL
# ─────────────────────────────────────────────
COLORS = {
    'header_dark':   '1F3864',
    'header_blue':   '2E75B6',
    'group_header':  'D9E1F2',  # секция Основные/Сопутствующие/Доставка
    'total_section': 'BDD7EE',  # ИТОГО по секции
    'grand_total':   '1F3864',  # ИТОГО всего
    'row_white':     'FFFFFF',
    'ozon_blue':     '005BFF',
    'wb_violet':     '7B2FBE',
    'yellow':        'FFEB9C',
}

def hfill(color): return PatternFill('solid', fgColor=color)
def hfont(bold=False, white=False, size=10):
    return Font(bold=bold, color='FFFFFF' if white else '000000', size=size, name='Calibri')
def center(): return Alignment(horizontal='center', vertical='center', wrap_text=True)
def left():   return Alignment(horizontal='left',   vertical='center', wrap_text=True)
def right():  return Alignment(horizontal='right',  vertical='center')

BORDER_THIN = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'),  bottom=Side(style='thin')
)

def _set_col_widths(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def _write_cell(ws, row, col, value, fill_color, bold=False, white=False,
                fmt=None, align=None, size=10):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = hfill(fill_color)
    c.font = hfont(bold=bold, white=white, size=size)
    c.border = BORDER_THIN
    c.alignment = align or (left() if col == 1 else right())
    if fmt:
        c.number_format = fmt
    return c


# ─────────────────────────────────────────────
#  ГРУППИРОВКА ДАННЫХ КАНАЛА
# ─────────────────────────────────────────────

def group_channel_df(df: pd.DataFrame, sebes_dict: dict) -> dict:
    """
    Разбивает DataFrame канала на группы: основные, сопутствующие, доставка.
    Возвращает dict с ключами 'основные', 'сопутствующие', 'доставка'.
    Каждый ключ — список dict строк с полями:
      name, qty, rev, sebes1, sebes_all, profit, margin
    """
    groups = {'основные': [], 'сопутствующие': [], 'доставка': []}

    for _, r in df.iterrows():
        name = str(r.get('Наименование', '')).strip()
        if not name or name == 'nan':
            continue

        qty_f     = _to_float(r.get('Кол-во'))
        rev       = _to_float(r.get('Сумма продаж'))
        profit    = _to_float(r.get('Прибыль'))
        sebes_auto = r.get('Себес_авто')
        sebes_1c   = _to_float(r.get('Себес_1С'))
        # Себестоимость 1 ед: сначала авто, потом 1С
        sebes1 = float(sebes_auto) if sebes_auto else sebes_1c
        sebes_all = sebes1 * qty_f if qty_f else 0
        margin = (profit / rev) if rev else 0

        cat = classify_item(name)

        # Для доставки себестоимость = 0
        if cat == 'доставка':
            sebes1 = 0
            sebes_all = 0

        groups[cat].append({
            'name':      name,
            'qty':       qty_f,
            'rev':       rev,
            'sebes1':    sebes1,
            'sebes_all': sebes_all,
            'profit':    profit,
            'margin':    margin,
        })

    return groups


def calc_group_totals(rows: list) -> dict:
    """Считает ИТОГО по группе строк."""
    qty   = sum(r['qty'] or 0 for r in rows)
    rev   = sum(r['rev'] or 0 for r in rows)
    sebes = sum(r['sebes_all'] or 0 for r in rows)
    prof  = sum(r['profit'] or 0 for r in rows)
    marg  = (prof / rev) if rev else 0
    return {'qty': qty, 'rev': rev, 'sebes_all': sebes, 'profit': prof, 'margin': marg}


# ─────────────────────────────────────────────
#  ПОСТРОЕНИЕ EXCEL-ОТЧЁТА
# ─────────────────────────────────────────────

def build_excel(ms_data: dict, ozon_data: dict, wb_data: dict,
                period: str, sebes_dict: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    ws_svod = wb.create_sheet('СВОДНЫЙ')

    # Сначала строим листы каналов и собираем итоги для сводного
    channel_display = [
        ('Интернет_сен26',  'Интернет',  'SP.Candle интернет магазин'),
        ('ОПТ_сен26',       'ОПТ',       'ОПТ'),
        ('СТМ_сен26',       'СТМ',       'СЗ/СТМ'),
        ('Столярный_сен26', 'Столярный', 'Магазин Столярный 9'),
    ]

    svod_rows = []  # данные для сводного листа
    for key, sheet_name, display_name in channel_display:
        if key not in ms_data:
            continue
        df = ms_data[key]
        groups = group_channel_df(df, sebes_dict)

        ws_ch = wb.create_sheet(f'МС_{sheet_name}')
        _build_channel_sheet(ws_ch, groups, display_name, period)

        # Итоги для сводного (основные + сопутствующие, без доставки)
        main_rows  = groups['основные'] + groups['сопутствующие']
        deliv_rows = groups['доставка']
        t_main  = calc_group_totals(main_rows)
        t_deliv = calc_group_totals(deliv_rows)
        svod_rows.append({
            'name':        display_name,
            'main':        t_main,
            'delivery':    t_deliv,
        })

    ozon_real = {k: v for k, v in ozon_data.items() if not k.startswith('_')}

    _build_svod_sheet(ws_svod, svod_rows, ozon_real, wb_data, period)

    if ozon_real:
        ws_ozon_svod = wb.create_sheet('Ozon_Сводка')
        _build_ozon_svod(ws_ozon_svod, ozon_real, period)

        if 'детализация' in ozon_real:
            ws_ozon_det = wb.create_sheet('Ozon_Детализация')
            _build_ozon_detail(ws_ozon_det, ozon_real['детализация'], sebes_dict)

        if 'по_типам' in ozon_real:
            ws_ozon_types = wb.create_sheet('Ozon_По типам')
            _build_raw_sheet(ws_ozon_types, ozon_real['по_типам'],
                             'Продажи Ozon по типам товара', COLORS['ozon_blue'])

    if wb_data and 'детализация' in wb_data:
        ws_wb_svod = wb.create_sheet('WB_Сводка')
        _build_wb_svod(ws_wb_svod, wb_data['детализация'], period, sebes_dict)

        ws_wb_det = wb.create_sheet('WB_Детализация')
        _build_wb_detail(ws_wb_det, wb_data['детализация'], sebes_dict)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ─────────────────────────────────────────────
#  ЛИСТ: КАНАЛ (Интернет / ОПТ / СТМ / Столярный)
# ─────────────────────────────────────────────

def _build_channel_sheet(ws, groups: dict, channel_name: str, period: str):
    ws.sheet_view.showGridLines = False
    COLS = ['A', 'B', 'C', 'D', 'E', 'F', 'G']
    NCOLS = 7

    # Заголовок
    ws.merge_cells('A1:G1')
    c = ws['A1']
    c.value = f'Сводный отчёт — {channel_name} — {period}'
    c.fill = hfill(COLORS['header_blue'])
    c.font = hfont(bold=True, white=True, size=12)
    c.alignment = center()
    ws.row_dimensions[1].height = 25

    # Шапка таблицы
    hdrs = ['Наименование', 'Кол-во', 'Выручка', 'Себес 1 ед', 'Себес итого', 'Прибыль', 'Маржа %']
    row = 2
    for ci, h in enumerate(hdrs, 1):
        _write_cell(ws, row, ci, h, 'BDD7EE', bold=True, align=center())
    row += 1

    section_labels = [
        ('основные',       'Основные товары'),
        ('сопутствующие',  'Сопутствующие товары'),
        ('доставка',       'Доставка и транспорт'),
    ]

    grand_qty = grand_rev = grand_sebes = grand_prof = 0.0

    for cat_key, cat_label in section_labels:
        items = groups.get(cat_key, [])

        # Заголовок секции (синяя строка)
        ws.merge_cells(f'A{row}:G{row}')
        c = ws.cell(row=row, column=1, value=cat_label)
        c.fill = hfill(COLORS['header_blue'])
        c.font = hfont(bold=True, white=True)
        c.alignment = left()
        c.border = BORDER_THIN
        ws.row_dimensions[row].height = 16
        row += 1

        if not items:
            row += 1
            continue

        for r in items:
            sebes1_out = r['sebes1'] if r['sebes1'] else None
            vals = [r['name'], r['qty'], r['rev'], sebes1_out,
                    r['sebes_all'] or None, r['profit'], r['margin']]
            fmts = [None, '#,##0', '#,##0.00', '#,##0.00', '#,##0.00', '#,##0.00', '0.0%']
            for ci, (v, fmt) in enumerate(zip(vals, fmts), 1):
                _write_cell(ws, row, ci, v, COLORS['row_white'], fmt=fmt)
            row += 1

        # ИТОГО секции
        t = calc_group_totals(items)
        label = f'ИТОГО {cat_label.lower()}'
        vals_t = [label, t['qty'], t['rev'], None, t['sebes_all'], t['profit'], t['margin']]
        fmts_t = [None, '#,##0', '#,##0.00', None, '#,##0.00', '#,##0.00', '0.0%']
        for ci, (v, fmt) in enumerate(zip(vals_t, fmts_t), 1):
            _write_cell(ws, row, ci, v, COLORS['total_section'], bold=True, fmt=fmt)
        row += 1

        if cat_key != 'доставка':
            grand_qty   += t['qty']
            grand_rev   += t['rev']
            grand_sebes += t['sebes_all']
            grand_prof  += t['profit']

    # Итоговая строка ИТОГО (тёмно-синяя)
    grand_margin = (grand_prof / grand_rev) if grand_rev else 0
    vals_g = ['ИТОГО', grand_qty, grand_rev, None, grand_sebes, grand_prof, grand_margin]
    fmts_g = [None, '#,##0', '#,##0.00', None, '#,##0.00', '#,##0.00', '0.0%']
    for ci, (v, fmt) in enumerate(zip(vals_g, fmts_g), 1):
        _write_cell(ws, row, ci, v, COLORS['grand_total'],
                    bold=True, white=True, fmt=fmt)

    _set_col_widths(ws, {'A': 38, 'B': 10, 'C': 16, 'D': 16, 'E': 16, 'F': 14, 'G': 10})


# ─────────────────────────────────────────────
#  ЛИСТ: СВОДНЫЙ
# ─────────────────────────────────────────────

def _build_svod_sheet(ws, svod_rows: list, ozon_data: dict, wb_data: dict, period: str):
    ws.sheet_view.showGridLines = False

    # Заголовок
    ws.merge_cells('A1:H1')
    c = ws['A1']
    c.value = f'СВОДНЫЙ ОТЧЁТ SP CANDLE — {period.upper()}'
    c.fill = hfill(COLORS['header_dark'])
    c.font = hfont(bold=True, white=True, size=13)
    c.alignment = center()
    ws.row_dimensions[1].height = 30

    row = 3

    # ── Блок 1С ──
    ws.merge_cells(f'A{row}:H{row}')
    c = ws[f'A{row}']
    c.value = '1С — ПРОДАЖИ ПО КАНАЛАМ'
    c.fill = hfill(COLORS['header_blue'])
    c.font = hfont(bold=True, white=True)
    c.alignment = center()
    row += 1

    hdrs = ['Канал', 'Кол-во', 'Выручка', 'Себес итого', 'Прибыль', 'Маржа %']
    for ci, h in enumerate(hdrs, 1):
        _write_cell(ws, row, ci, h, 'BDD7EE', bold=True, align=center())
    row += 1

    # Секция основных + сопутствующих товаров
    ws.merge_cells(f'A{row}:F{row}')
    c = ws.cell(row=row, column=1, value='Основные и сопутствующие товары')
    c.fill = hfill(COLORS['group_header'])
    c.font = hfont(bold=True)
    c.alignment = left()
    c.border = BORDER_THIN
    row += 1

    tot_qty = tot_rev = tot_sebes = tot_prof = 0.0
    for sr in svod_rows:
        t = sr['main']
        marg = (t['profit'] / t['rev']) if t['rev'] else 0
        vals = [sr['name'], t['qty'], t['rev'], t['sebes_all'], t['profit'], marg]
        fmts = [None, '#,##0', '#,##0.00', '#,##0.00', '#,##0.00', '0.0%']
        for ci, (v, fmt) in enumerate(zip(vals, fmts), 1):
            _write_cell(ws, row, ci, v, COLORS['row_white'], fmt=fmt)
        tot_qty   += t['qty']
        tot_rev   += t['rev']
        tot_sebes += t['sebes_all']
        tot_prof  += t['profit']
        row += 1

    t_marg = (tot_prof / tot_rev) if tot_rev else 0
    vals_t = ['ИТОГО', tot_qty, tot_rev, tot_sebes, tot_prof, t_marg]
    fmts_t = [None, '#,##0', '#,##0.00', '#,##0.00', '#,##0.00', '0.0%']
    for ci, (v, fmt) in enumerate(zip(vals_t, fmts_t), 1):
        _write_cell(ws, row, ci, v, COLORS['total_section'], bold=True, fmt=fmt)
    row += 1

    # Секция доставки
    ws.merge_cells(f'A{row}:F{row}')
    c = ws.cell(row=row, column=1, value='Доставка и транспорт')
    c.fill = hfill(COLORS['group_header'])
    c.font = hfont(bold=True)
    c.alignment = left()
    c.border = BORDER_THIN
    row += 1

    del_qty = del_rev = 0.0
    for sr in svod_rows:
        t = sr['delivery']
        vals = [sr['name'], t['qty'], t['rev'], 0, t['rev'], '1']
        fmts = [None, '#,##0', '#,##0.00', None, '#,##0.00', '0%']
        for ci, (v, fmt) in enumerate(zip([sr['name'], t['qty'], t['rev']], [None, '#,##0', '#,##0.00']), 1):
            _write_cell(ws, row, ci, v, COLORS['row_white'], fmt=fmt)
        del_qty += t['qty']
        del_rev += t['rev']
        row += 1

    vals_d = ['ИТОГО', del_qty, del_rev]
    for ci, (v, fmt) in enumerate(zip(vals_d, [None, '#,##0', '#,##0.00']), 1):
        _write_cell(ws, row, ci, v, COLORS['total_section'], bold=True, fmt=fmt)
    row += 2

    # ── Блок Ozon ──
    ws.merge_cells(f'A{row}:H{row}')
    c = ws[f'A{row}']
    c.value = 'OZON — НАЧИСЛЕНИЯ'
    c.fill = hfill(COLORS['ozon_blue'])
    c.font = hfont(bold=True, white=True)
    c.alignment = center()
    row += 1

    for ci, h in enumerate(['Статья', 'Сумма, руб.'], 1):
        _write_cell(ws, row, ci, h, 'BDD7EE', bold=True, align=center())
    row += 1

    ozon_summary = _extract_ozon_summary(ozon_data)
    for stat, val in ozon_summary.items():
        is_total = any(x in stat.lower() for x in ('итого', 'перечислению', 'начислено'))
        fill = COLORS['total_section'] if is_total else COLORS['row_white']
        _write_cell(ws, row, 1, stat, fill, bold=is_total)
        _write_cell(ws, row, 2, val, fill, bold=is_total, fmt='#,##0.00')
        row += 1
    row += 1

    # ── Блок WB ──
    ws.merge_cells(f'A{row}:H{row}')
    c = ws[f'A{row}']
    c.value = 'WILDBERRIES — НАЧИСЛЕНИЯ'
    c.fill = hfill(COLORS['wb_violet'])
    c.font = hfont(bold=True, white=True)
    c.alignment = center()
    row += 1

    for ci, h in enumerate(['Статья', 'Сумма, руб.'], 1):
        _write_cell(ws, row, ci, h, 'BDD7EE', bold=True, align=center())
    row += 1

    wb_summary = _extract_wb_summary(wb_data)
    for stat, val in wb_summary.items():
        is_total = 'итого' in stat.lower()
        fill = COLORS['total_section'] if is_total else COLORS['row_white']
        _write_cell(ws, row, 1, stat, fill, bold=is_total)
        _write_cell(ws, row, 2, val, fill, bold=is_total, fmt='#,##0.00')
        row += 1

    _set_col_widths(ws, {'A': 38, 'B': 14, 'C': 16, 'D': 16, 'E': 14, 'F': 10, 'G': 12, 'H': 12})


# ─────────────────────────────────────────────
#  ЛИСТ: OZON
# ─────────────────────────────────────────────

def _extract_ozon_summary(ozon_data: dict) -> dict:
    if not ozon_data or 'сводка' not in ozon_data:
        return {}
    df = ozon_data['сводка']
    result = {}
    for i, r in df.iterrows():
        label = str(r.iloc[0]).strip() if pd.notna(r.iloc[0]) else ''
        val   = r.iloc[1] if len(r) > 1 and pd.notna(r.iloc[1]) else None
        if label and label != 'nan' and val is not None:
            try:
                result[label] = float(val)
            except (TypeError, ValueError):
                pass
    return result


def _build_ozon_svod(ws, ozon_data: dict, period: str):
    ws.sheet_view.showGridLines = False
    ws.merge_cells('A1:C1')
    c = ws['A1']
    c.value = f'OZON — СВОДКА НАЧИСЛЕНИЙ — {period.upper()}'
    c.fill = hfill(COLORS['ozon_blue'])
    c.font = hfont(bold=True, white=True, size=12)
    c.alignment = center()
    ws.row_dimensions[1].height = 25

    summary = _extract_ozon_summary(ozon_data)
    row = 3
    for ci, h in enumerate(['Статья', 'Сумма, руб.'], 1):
        _write_cell(ws, row, ci, h, 'BDD7EE', bold=True, align=center())
    row += 1

    for stat, val in summary.items():
        is_total = any(x in stat.lower() for x in ('итого', 'перечислению', 'начислено'))
        fill = COLORS['total_section'] if is_total else COLORS['row_white']
        _write_cell(ws, row, 1, stat, fill, bold=is_total)
        _write_cell(ws, row, 2, val, fill, bold=is_total, fmt='#,##0.00')
        row += 1

    if 'по_типам' in ozon_data:
        row += 1
        df = ozon_data['по_типам']
        ws.merge_cells(f'A{row}:F{row}')
        c = ws[f'A{row}']
        c.value = 'ПРОДАЖИ ПО ТИПАМ ТОВАРА (Ozon)'
        c.fill = hfill(COLORS['ozon_blue'])
        c.font = hfont(bold=True, white=True)
        c.alignment = center()
        row += 1
        _dump_df_to_ws(ws, df, row)

    _set_col_widths(ws, {'A': 45, 'B': 16, 'C': 14, 'D': 14, 'E': 14, 'F': 16})


def _build_ozon_detail(ws, df: pd.DataFrame, sebes_dict: dict):
    ws.sheet_view.showGridLines = False
    ws.merge_cells('A1:J1')
    c = ws['A1']
    c.value = 'OZON — ДЕТАЛИЗАЦИЯ НАЧИСЛЕНИЙ'
    c.fill = hfill(COLORS['ozon_blue'])
    c.font = hfont(bold=True, white=True, size=12)
    c.alignment = center()
    ws.row_dimensions[1].height = 25

    if 'Название товара' in df.columns:
        df = df.copy()
        df['Себес (авто)'] = df['Название товара'].apply(lambda x: get_sebes(x, sebes_dict))

    row = 3
    for ci, h in enumerate(df.columns, 1):
        _write_cell(ws, row, ci, str(h), COLORS['ozon_blue'], bold=True, white=True, align=center())
    row += 1

    for _, dr in df.iterrows():
        for ci, v in enumerate(dr.values, 1):
            v_out = None if pd.isna(v) else v
            _write_cell(ws, row, ci, v_out, COLORS['row_white'],
                        align=left() if ci <= 3 else right())
        row += 1

    _set_col_widths(ws, {'A': 18, 'B': 20, 'C': 22, 'D': 25, 'E': 14, 'F': 14,
                         'G': 40, 'H': 10, 'I': 12, 'J': 16})


# ─────────────────────────────────────────────
#  ЛИСТ: WB
# ─────────────────────────────────────────────

def _extract_wb_summary(wb_data: dict) -> dict:
    if not wb_data or 'детализация' not in wb_data:
        return {}
    df = wb_data['детализация']
    result = {}

    money_col    = 'К перечислению Продавцу за реализованный Товар'
    delivery_col = 'Услуги по доставке товара покупателю'
    storage_col  = 'Хранение'
    fine_col     = 'Общая сумма штрафов'
    logistic_col = 'Возмещение издержек по перемещению и операционной обработке товара'
    obosnov      = 'Обоснование для оплаты'

    mask_sale = df.get(obosnov, pd.Series(dtype=str)) == 'Продажа'
    if money_col in df.columns:
        result['К перечислению за товар (Продажи)'] = float(df.loc[mask_sale, money_col].sum()) if mask_sale.any() else 0.0
    if delivery_col in df.columns:
        result['Услуги доставки'] = float(df[delivery_col].sum())
    if storage_col in df.columns:
        result['Хранение'] = float(df[storage_col].sum())
    if logistic_col in df.columns:
        result['Логистика (перемещение)'] = float(df[logistic_col].sum())
    if fine_col in df.columns:
        result['Штрафы'] = float(df[fine_col].sum())
    if money_col in df.columns:
        result['ИТОГО К ПЕРЕЧИСЛЕНИЮ'] = float(df[money_col].sum())

    return result


def _build_wb_svod(ws, df: pd.DataFrame, period: str, sebes_dict: dict):
    ws.sheet_view.showGridLines = False
    ws.merge_cells('A1:C1')
    c = ws['A1']
    c.value = f'WILDBERRIES — СВОДКА — {period.upper()}'
    c.fill = hfill(COLORS['wb_violet'])
    c.font = hfont(bold=True, white=True, size=12)
    c.alignment = center()

    summary = _extract_wb_summary({'детализация': df})
    row = 3
    for ci, h in enumerate(['Статья', 'Сумма, руб.'], 1):
        _write_cell(ws, row, ci, h, 'BDD7EE', bold=True, align=center())
    row += 1

    for stat, val in summary.items():
        is_total = 'итого' in stat.lower()
        fill = COLORS['total_section'] if is_total else COLORS['row_white']
        _write_cell(ws, row, 1, stat, fill, bold=is_total)
        _write_cell(ws, row, 2, val, fill, bold=is_total, fmt='#,##0.00')
        row += 1

    _set_col_widths(ws, {'A': 45, 'B': 16})


def _build_wb_detail(ws, df: pd.DataFrame, sebes_dict: dict):
    ws.sheet_view.showGridLines = False
    ws.merge_cells('A1:J1')
    c = ws['A1']
    c.value = 'WILDBERRIES — ДЕТАЛИЗАЦИЯ (все еженедельные отчёты)'
    c.fill = hfill(COLORS['wb_violet'])
    c.font = hfont(bold=True, white=True, size=12)
    c.alignment = center()

    key_cols = [
        'Название', 'Артикул поставщика', 'Тип документа', 'Обоснование для оплаты',
        'Дата заказа покупателем', 'Дата продажи', 'Кол-во',
        'Вайлдберриз реализовал Товар (Пр)',
        'К перечислению Продавцу за реализованный Товар',
        'Услуги по доставке товара покупателю',
        'Хранение', 'Общая сумма штрафов',
        'Возмещение издержек по перемещению и операционной обработке товара',
    ]
    available = [col for col in key_cols if col in df.columns]
    df_out = df[available].copy()
    if 'Название' in df_out.columns:
        df_out['Себес (авто)'] = df_out['Название'].apply(lambda x: get_sebes(x, sebes_dict))

    row = 3
    for ci, h in enumerate(df_out.columns, 1):
        _write_cell(ws, row, ci, str(h), COLORS['wb_violet'], bold=True, white=True, align=center())
    row += 1

    for _, dr in df_out.iterrows():
        for ci, v in enumerate(dr.values, 1):
            v_out = None if (isinstance(v, float) and np.isnan(v)) else v
            _write_cell(ws, row, ci, v_out, COLORS['row_white'],
                        align=left() if ci <= 2 else right())
        row += 1

    _set_col_widths(ws, {'A': 40, 'B': 14, 'C': 14, 'D': 30, 'E': 18, 'F': 18,
                         'G': 8, 'H': 14, 'I': 16, 'J': 14, 'K': 12, 'L': 12, 'M': 16})


# ─────────────────────────────────────────────
#  ВСПОМОГАТЕЛЬНЫЕ
# ─────────────────────────────────────────────

def _build_raw_sheet(ws, df_raw: pd.DataFrame, title: str, color: str):
    ws.sheet_view.showGridLines = False
    max_col = max(df_raw.shape[1], 1)
    end_col = get_column_letter(max_col)
    ws.merge_cells(f'A1:{end_col}1')
    c = ws['A1']
    c.value = title
    c.fill = hfill(color)
    c.font = hfont(bold=True, white=True, size=11)
    c.alignment = center()
    _dump_df_to_ws(ws, df_raw, start_row=3)


def _dump_df_to_ws(ws, df: pd.DataFrame, start_row: int):
    for i, (_, r) in enumerate(df.iterrows()):
        for ci, v in enumerate(r.values, 1):
            v_out = None if (isinstance(v, float) and np.isnan(v)) else v
            c = ws.cell(row=start_row + i, column=ci, value=v_out)
            c.fill = hfill(COLORS['row_white'])
            c.border = BORDER_THIN
            c.alignment = left() if ci == 1 else right()


# ─────────────────────────────────────────────
#  STREAMLIT UI
# ─────────────────────────────────────────────

def render_sebes_editor():
    st.subheader('⚙️ Настройки себестоимости')
    st.caption('Здесь можно изменить себестоимость для любого типа товара.')

    sebes = load_sebes()

    _, col_reset = st.columns([3, 1])
    with col_reset:
        if st.button('↺ Сбросить к исходным', use_container_width=True):
            save_sebes(dict(DEFAULT_SEBES))
            st.success('Сброшено!')
            st.rerun()

    df_sebes = pd.DataFrame(
        [{'Тип товара (подстрока в названии)': k, 'Себестоимость, руб.': v}
         for k, v in sebes.items()]
    )

    edited = st.data_editor(
        df_sebes,
        num_rows='dynamic',
        use_container_width=True,
        column_config={
            'Тип товара (подстрока в названии)': st.column_config.TextColumn(width='large'),
            'Себестоимость, руб.': st.column_config.NumberColumn(min_value=0, format='%.2f', width='medium'),
        },
        hide_index=True,
        key='sebes_editor',
    )

    if st.button('💾 Сохранить себестоимость', type='primary', use_container_width=True):
        new_sebes = {}
        for _, row in edited.iterrows():
            key = str(row['Тип товара (подстрока в названии)']).strip().lower()
            val = row['Себестоимость, руб.']
            if key and not pd.isna(val):
                new_sebes[key] = float(val)
        save_sebes(new_sebes)
        st.success(f'✅ Сохранено {len(new_sebes)} позиций!')
        st.rerun()

    return sebes


def main():
    st.set_page_config(
        page_title='SP CANDLE — Отчёт',
        page_icon='🕯️',
        layout='wide',
    )

    st.title('🕯️ SP CANDLE — Сводный отчёт за месяц')

    tab_main, tab_sebes = st.tabs(['📊 Формирование отчёта', '⚙️ Себестоимость'])

    with tab_sebes:
        sebes_dict = render_sebes_editor()

    with tab_main:
        st.markdown('Загрузите файлы из 1С, Ozon и Wildberries — получите единый Excel-отчёт.')

        col1, col2 = st.columns([1, 1])
        months_ru = ['Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь',
                     'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь']
        with col1:
            month_num = st.selectbox('Месяц отчёта', range(1, 13),
                                     index=datetime.now().month - 2,
                                     format_func=lambda x: months_ru[x - 1])
        with col2:
            year = st.number_input('Год', value=datetime.now().year, step=1, min_value=2020)
        period = f'{months_ru[month_num - 1]} {year}'

        st.divider()

        with st.expander('📅 Период фильтрации WB-отчётов', expanded=False):
            st.caption('WB-отчёты содержат данные за несколько месяцев. Укажите период.')
            import datetime as dt, calendar
            wb_col1, wb_col2 = st.columns(2)
            with wb_col1:
                default_from = dt.date(int(year), month_num, 1)
                wb_date_from = st.date_input('Дата С', value=default_from, key='wb_from')
            with wb_col2:
                last_day = calendar.monthrange(int(year), month_num)[1]
                default_to = dt.date(int(year), month_num, last_day)
                wb_date_to = st.date_input('Дата ПО', value=default_to, key='wb_to')

        st.divider()

        st.subheader('📂 1С — Прибыльность по товарам (4 XLS-файла)')
        st.caption('Каждый файл — отдельный канал: Интернет, ОПТ, СТМ, Столярный')
        col_a, col_b = st.columns(2)
        with col_a:
            f_internet = st.file_uploader('Интернет магазин', type=['xls'], key='internet')
            f_opt      = st.file_uploader('ОПТ', type=['xls'], key='opt')
        with col_b:
            f_stm      = st.file_uploader('СЗ/СТМ', type=['xls'], key='stm')
            f_stolyar  = st.file_uploader('Магазин Столярный 9', type=['xls'], key='stolyar')

        st.subheader('🔵 Ozon — отчёт начислений (xlsx)')
        f_ozon = st.file_uploader('Отчёт Ozon', type=['xlsx'], key='ozon')

        st.subheader('🟣 Wildberries — еженедельные отчёты (zip-архивы)')
        st.caption('Загрузите все еженедельные отчёты — лишние месяцы будут отфильтрованы')
        f_wb_zips = st.file_uploader('WB отчёты', type=['zip'],
                                      accept_multiple_files=True, key='wb')

        st.divider()

        if st.button('🚀 Сформировать отчёт', type='primary', use_container_width=True):
            sebes_dict = load_sebes()

            with st.spinner('Обрабатываю данные...'):
                ms_data = {}
                channel_files = {
                    'Интернет_сен26':  f_internet,
                    'ОПТ_сен26':       f_opt,
                    'СТМ_сен26':       f_stm,
                    'Столярный_сен26': f_stolyar,
                }
                for key, f in channel_files.items():
                    if f is not None:
                        ch_name = key.split('_')[0]
                        df = parse_1c_file(f, ch_name, sebes_dict)
                        ms_data[key] = df
                        st.success(f'✅ {ch_name}: {len(df)} позиций')

                if not ms_data:
                    st.warning('⚠️ Не загружен ни один файл 1С')

                ozon_data = {}
                if f_ozon is not None:
                    try:
                        ozon_data = parse_ozon_file(f_ozon)
                        sheets = ozon_data.get('_sheet_names', [])
                        found  = [k for k in ozon_data if not k.startswith('_')]
                        st.success(f'✅ Ozon: листы в файле {sheets} → найдено: {found}')
                    except Exception as e:
                        import traceback
                        st.error(f'❌ Ошибка Ozon: {e}\n{traceback.format_exc()}')

                wb_data = {}
                if f_wb_zips:
                    wb_data = parse_wb_files(f_wb_zips, wb_date_from, wb_date_to)
                    n = len(wb_data.get('детализация', pd.DataFrame()))
                    st.success(f'✅ WB: {n} строк за выбранный период')

                if not ms_data and not ozon_data and not wb_data:
                    st.error('Загрузите хотя бы один файл!')
                else:
                    excel_bytes = build_excel(ms_data, ozon_data, wb_data, period, sebes_dict)
                    fname = f'SP_CANDLE_Отчет_{period.replace(" ", "_")}.xlsx'
                    st.download_button(
                        label='⬇️ Скачать Excel-отчёт',
                        data=excel_bytes,
                        file_name=fname,
                        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        type='primary',
                        use_container_width=True,
                    )
                    st.balloons()


if __name__ == '__main__':
    main()
